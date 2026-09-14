import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# Colores usados en el Excel
COLOR_ENCABEZADO = "2F5496"
COLOR_DUPLICADO = "FFF2CC"     # amarillo suave
COLOR_VENCIDA = "F8CBAD"       # rojo/salmón suave
COLOR_VENCE_MANANA = "FCE4D6"  # naranja suave
COLOR_PROXIMAMENTE = "D9E1F2"  # celeste suave


def generar_excel(resultado, archivo_salida, dividir_control=False):
    """
    Recibe el DataFrame ya cruzado (resultado) y genera el Excel final.

    Siempre incluye la hoja "Promociones en Stock" (el listado completo,
    con la marca de "Control"/duplicado si corresponde).

    Para la parte de Control, pensada para imprimir y que el repositor
    vaya tildando a medida que coloca cada cartel, hay dos modos:

    - dividir_control=False (default): una única hoja "Control" con
      todas las promociones juntas.
    - dividir_control=True: una hoja "Control - <nombre>" por cada
      hoja del Excel de marketing en la que aparecía la promoción
      (ej. "Control - CVMH", "Control - LAVANDINA LOREAL"). Si un
      mismo producto tenía promo en más de una hoja de marketing,
      aparece en cada una de esas hojas de Control.

    En ningún caso las hojas de Control muestran la leyenda de
    "duplicado": esa marca es solo para la hoja principal.
    """

    libro = openpyxl.Workbook()

    _armar_hoja_principal(libro, resultado)

    if dividir_control:
        _armar_hojas_control_divididas(libro, resultado)
    else:
        _armar_hoja_control(libro, "Control", resultado)

    libro.save(archivo_salida)


# ============================================================
# HOJA 1: PROMOCIONES EN STOCK (listado completo)
# ============================================================

def _armar_hoja_principal(libro, resultado):

    hoja = libro.active
    hoja.title = "Promociones en Stock"

    # Orden y nombres de columnas para esta hoja
    columnas = [
        ("Linea", "Línea"),
        ("Descripción", "Descripción"),
        ("Código de Barras", "Código de Barras"),
        ("Stock", "Stock"),
        ("DTO", "Promoción"),
        ("Precio de Venta", "Precio de Venta"),
        ("Inicio", "Inicio"),
        ("Fin", "Fin"),
        ("Estado", "Estado"),
    ]

    tiene_duplicados = "Revisar duplicado" in resultado.columns

    if tiene_duplicados:
        columnas.append(("Revisar duplicado", "Control"))

    encabezados = [nombre_mostrado for _, nombre_mostrado in columnas]
    hoja.append(encabezados)

    for _, fila in resultado.iterrows():
        hoja.append([fila[columna_original] for columna_original, _ in columnas])

    _dar_estilo_encabezado(hoja)

    # Ubicar columnas clave por nombre (por si el orden cambia en el futuro)
    posiciones = {nombre: i + 1 for i, (_, nombre) in enumerate(columnas)}

    _resaltar_por_estado(hoja, posiciones.get("Estado"))

    if tiene_duplicados:
        _resaltar_duplicados(hoja, posiciones.get("Control"))

    if "Precio de Venta" in posiciones:
        _formato_moneda(hoja, posiciones["Precio de Venta"])

    anchos = {
        "A": 18, "B": 40, "C": 16, "D": 8,
        "E": 12, "F": 15, "G": 13, "H": 13,
        "I": 14, "J": 12
    }
    for letra, ancho in anchos.items():
        hoja.column_dimensions[letra].width = ancho

    hoja.freeze_panes = "A2"

    _configurar_impresion(hoja)


# ============================================================
# HOJA 2: CONTROL (para imprimir y tildar en el local)
# ============================================================

def _armar_hoja_control(libro, nombre_hoja, datos):
    """
    Arma una hoja de Control (para imprimir y tildar) con los datos
    recibidos. Se usa tanto para la hoja única "Control" como para
    cada una de las hojas divididas por hoja de marketing.

    Nunca incluye la marca de "duplicado": esa leyenda es exclusiva
    de la hoja principal "Promociones en Stock".
    """

    hoja = libro.create_sheet(nombre_hoja)

    encabezados = [
        "Línea", "Descripción", "Código de Barras", "Stock",
        "Promoción", "Precio de Venta", "Inicio", "Fin", "Vigencia", "Control"
    ]
    hoja.append(encabezados)

    # Ordenado por Línea y Descripción, para que sea más fácil
    # recorrer la farmacia por marca/sector.
    datos_ordenados = datos.sort_values(["Linea", "Descripción"])

    for _, fila in datos_ordenados.iterrows():
        hoja.append([
            fila["Linea"],
            fila["Descripción"],
            fila["Código de Barras"],
            fila["Stock"],
            fila["DTO"],
            fila["Precio de Venta"],
            fila["Inicio"],
            fila["Fin"],
            fila["Estado"],
            ""  # columna en blanco para tildar a mano
        ])

    _dar_estilo_encabezado(hoja)

    columna_precio = 6
    columna_vigencia = 9
    columna_control = 10

    _resaltar_por_estado(hoja, columna_vigencia)
    _formato_moneda(hoja, columna_precio)

    # Cuadrícula completa en toda la tabla, para que sea más fácil
    # de seguir con la vista al imprimir y tildar a mano.
    borde = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    for fila in range(1, hoja.max_row + 1):
        for columna in range(1, hoja.max_column + 1):
            hoja.cell(fila, columna).border = borde

    for fila in range(2, hoja.max_row + 1):
        celda_control = hoja.cell(fila, columna_control)
        celda_control.alignment = Alignment(horizontal="center")
        hoja.row_dimensions[fila].height = 20

    anchos = {
        "A": 18, "B": 42, "C": 16, "D": 8,
        "E": 12, "F": 15, "G": 13, "H": 13, "I": 14, "J": 12
    }
    for letra, ancho in anchos.items():
        hoja.column_dimensions[letra].width = ancho

    hoja.freeze_panes = "A2"

    _configurar_impresion(hoja)


# Caracteres que Excel no permite en un nombre de hoja.
_CARACTERES_INVALIDOS_NOMBRE_HOJA = ["\\", "/", "*", "[", "]", ":", "?"]


def _nombre_hoja_control(nombre_hoja_marketing, nombres_ya_usados):
    """
    Arma el nombre "Control - <hoja de marketing>", sacando caracteres
    que Excel no permite en un nombre de hoja y respetando el límite
    de 31 caracteres. Si dos hojas de marketing generan el mismo
    nombre recortado, le agrega un número al final para no pisarlas.
    """

    nombre_base = str(nombre_hoja_marketing).strip() or "Sin nombre"

    for caracter in _CARACTERES_INVALIDOS_NOMBRE_HOJA:
        nombre_base = nombre_base.replace(caracter, "")

    nombre = f"Control - {nombre_base}"[:31]

    nombre_final = nombre
    contador = 2
    while nombre_final in nombres_ya_usados:
        sufijo = f" ({contador})"
        nombre_final = nombre[:31 - len(sufijo)] + sufijo
        contador += 1

    nombres_ya_usados.add(nombre_final)
    return nombre_final


def _armar_hojas_control_divididas(libro, resultado):
    """
    Genera una hoja "Control - <hoja de marketing>" por cada hoja del
    Excel de marketing que tuvo coincidencias en stock. Si un mismo
    producto tenía promoción en más de una hoja de marketing (caso de
    "duplicado"), aparece en cada una de esas hojas de Control.
    """

    nombres_ya_usados = set()

    hojas_marketing = sorted(
        resultado["Hoja"].dropna().unique(),
        key=lambda valor: str(valor).lower()
    )

    for hoja_marketing in hojas_marketing:
        datos_de_esta_hoja = resultado[resultado["Hoja"] == hoja_marketing]

        if datos_de_esta_hoja.empty:
            continue

        nombre_hoja_excel = _nombre_hoja_control(hoja_marketing, nombres_ya_usados)
        _armar_hoja_control(libro, nombre_hoja_excel, datos_de_esta_hoja)


# ============================================================
# FUNCIONES DE ESTILO REUTILIZABLES
# ============================================================

def _configurar_impresion(hoja):
    """
    Deja la hoja lista para imprimir en una impresora normal (A4):
    orientación horizontal, ajustada al ancho de la página, con
    la fila de encabezados repetida en cada página impresa.
    Funciona igual en Excel y en LibreOffice Calc.
    """

    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = "1:1"

    hoja.page_margins.left = 0.4
    hoja.page_margins.right = 0.4
    hoja.page_margins.top = 0.5
    hoja.page_margins.bottom = 0.5
    hoja.page_margins.header = 0.2
    hoja.page_margins.footer = 0.2


def _dar_estilo_encabezado(hoja):

    relleno = PatternFill(
        start_color=COLOR_ENCABEZADO,
        end_color=COLOR_ENCABEZADO,
        fill_type="solid"
    )

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center")


def _resaltar_por_estado(hoja, columna_estado):

    if columna_estado is None:
        return

    relleno_vencida = PatternFill(
        start_color=COLOR_VENCIDA, end_color=COLOR_VENCIDA, fill_type="solid"
    )
    relleno_manana = PatternFill(
        start_color=COLOR_VENCE_MANANA, end_color=COLOR_VENCE_MANANA, fill_type="solid"
    )
    relleno_proximamente = PatternFill(
        start_color=COLOR_PROXIMAMENTE, end_color=COLOR_PROXIMAMENTE, fill_type="solid"
    )

    for fila in range(2, hoja.max_row + 1):

        valor = hoja.cell(fila, columna_estado).value

        if valor == "VENCIDA":
            relleno = relleno_vencida
        elif valor == "VENCE MAÑANA":
            relleno = relleno_manana
        elif valor == "PRÓXIMAMENTE":
            relleno = relleno_proximamente
        else:
            continue

        for columna in range(1, hoja.max_column + 1):
            hoja.cell(fila, columna).fill = relleno


def _resaltar_duplicados(hoja, columna_duplicado):

    if columna_duplicado is None:
        return

    relleno = PatternFill(
        start_color=COLOR_DUPLICADO, end_color=COLOR_DUPLICADO, fill_type="solid"
    )

    for fila in range(2, hoja.max_row + 1):
        if hoja.cell(fila, columna_duplicado).value == "SI":
            # Solo pintamos la celda de "Revisar duplicado" para no tapar
            # el color de Estado (vencida/vence mañana), que es más importante.
            hoja.cell(fila, columna_duplicado).fill = relleno


def _formato_moneda(hoja, columna_precio):

    for fila in range(2, hoja.max_row + 1):
        hoja.cell(fila, columna_precio).number_format = '$#,##0.00'
